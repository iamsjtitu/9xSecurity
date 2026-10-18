"""Iteration 44: range report (any dates, PDF + Excel), plate search, Plate Recognizer usage card."""
import os
import sys
from datetime import datetime

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
import database  # noqa: E402
import report  # noqa: E402


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "REPORT_DIR", str(tmp_path / "reports"))
    d = database.EventDB(db_path=str(tmp_path / "e.db"))
    eid = d.add_event("truck", "Entry", "MH12AB1234", "/s/1.jpg", ts=datetime(2026, 10, 1, 9), plate_status="done", gate="Main")
    d.update_event_plate(eid, "MH12AB1234", source="cloud", status="done")
    d.add_event("truck", "Exit", "MH12AB1234", "/s/2.jpg", ts=datetime(2026, 10, 1, 18), plate_status="done", gate="Main")
    d.add_event("car", "Entry", "", "/s/3.jpg", ts=datetime(2026, 10, 15, 10), gate="Main")
    d.add_event("motorcycle", "Entry", "", "/s/4.jpg", ts=datetime(2026, 10, 31, 23, 59), gate="Back")
    d.add_event("person", "Exit", "", "/s/5.jpg", ts=datetime(2026, 10, 31, 23, 59, 30), count=2, gate="Back")
    d.add_event("truck", "Entry", "KA01MJ4567", "/s/6.jpg", ts=datetime(2026, 11, 1, 0, 0), plate_status="done", gate="Main")  # outside October
    return d


def test_range_window_validation():
    s, e = report.range_window("2026-10-01", "2026-10-31")
    assert (s, e) == (datetime(2026, 10, 1), datetime(2026, 11, 1))
    with pytest.raises(ValueError):
        report.range_window("2026-10-31", "2026-10-01")
    with pytest.raises(ValueError):
        report.range_window("2025-01-01", "2026-02-01")
    with pytest.raises(ValueError):
        report.range_window("2026-13-01", "2026-10-01")


def test_generate_range_month_pdf(db):
    summary, days, path = report.generate_range(db, "2026-10-01", "2026-10-31", "pdf")
    assert summary["events"] == 5 and summary["entry"] == 3 and summary["exit"] == 3  # person x2 exit counts members
    assert summary["total"]["truck"] == {"Entry": 1, "Exit": 1} and summary["gates"]["Back"]["person"]["Exit"] == 2
    assert len(days) == 31 and days[0]["date"] == "2026-10-01" and days[0]["entry"] == 1 and days[0]["exit"] == 1
    assert days[30]["events"] == 2 and days[14]["total"]["car"]["Entry"] == 1 and days[1]["events"] == 0
    assert report.period_text(summary) == "01-10-2026 se 31-10-2026 tak (31 din)"
    assert path.endswith("9x_Report_2026-10-01_to_2026-10-31.pdf") and open(path, "rb").read(5) == b"%PDF-"
    import fitz

    text = "\n".join(p.get_text() for p in fitz.open(path))
    assert "9x Security - Report" in text and "Day-wise (31 days)" in text and "All events (5)" in text
    assert "MH12AB1234" in text and "KA01MJ4567" not in text and "01-10 09:00 AM" in text  # date shown in multi-day lists


def test_generate_range_single_day_is_daily(db):
    summary, days, path = report.generate_range(db, "2026-10-01", "2026-10-01", "pdf")
    assert len(days) == 1 and summary["events"] == 2 and "01-10-2026 (Thursday) (poora din)" == report.period_text(summary)
    import fitz

    text = fitz.open(path)[0].get_text()
    assert "Daily Report" in text


def test_generate_range_xlsx(db):
    summary, days, path = report.generate_range(db, "2026-10-01", "2026-10-31", "xlsx")
    assert path.endswith(".xlsx") and os.path.getsize(path) > 4000
    from openpyxl import load_workbook

    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Day-wise", "Events"]
    ws = wb["Summary"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0][0] == "9x Security - Report" and "01-10-2026 - 31-10-2026 (31 days)" in rows[1][0]
    assert ["Gate", "Type", "Entry", "Exit", "Total"] in rows
    assert ["All gates", "Truck", 1, 1, 2] in rows and ["Back", "Person (paidal)", 0, 2, 2] in rows
    assert ["Main", "TOTAL", 2, 1, 3] in rows
    wd = wb["Day-wise"]
    assert wd.max_row == 32 and wd["A2"].value == "01-10-2026" and wd["B2"].value == 1 and wd["D2"].value == 1
    we = wb["Events"]
    assert we.max_row == 6 and we["G2"].value == "MH12AB1234" and we["H2"].value == "cloud" and we["I2"].value == "/s/1.jpg"
    assert we["G4"].value == "Not detected" and we["G5"].value == "-"  # car without plate / two-wheeler
    assert we.freeze_panes == "A2"


def test_generate_range_empty(db):
    summary, days, path = report.generate_range(db, "2026-09-01", "2026-09-07", "xlsx")
    assert summary["events"] == 0 and len(days) == 7 and os.path.exists(path)
    summary, days, path = report.generate_range(db, "2026-09-01", "2026-09-07", "pdf")
    import fitz

    assert "No Entry/Exit recorded in this period" in fitz.open(path)[0].get_text()


def test_plate_search_query_uses_like(db):
    rows = db.get_events(plate_filter="1234", limit=1000)
    assert {r["plate"] for r in rows} == {"MH12AB1234"} and len(rows) == 2
    assert db.get_events(plate_filter="MJ45")[0]["plate"] == "KA01MJ4567"
    assert db.get_events(plate_filter="ZZZZ") == []
