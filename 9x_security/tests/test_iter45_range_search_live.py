"""Iteration 45 (T1): live API checks for /api/report/range(.pdf/.xlsx), /api/plate/search,
/api/plate/cloud/usage (FAKE key only) + extra edge cases (month boundary, 366-day limit, fuzzy
ordering, xlsx column widths). Runs against the service on 127.0.0.1:8971."""
import io
import os
import sqlite3
import sys
import time
from datetime import datetime, timedelta

import pytest
import requests

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)
import database  # noqa: E402
import report  # noqa: E402

BASE = "http://127.0.0.1:8971"
DB = os.path.join(ROOT, "events.db")
SNAP = os.path.join(ROOT, "snapshots")
TAG = "TEST_iter45"
NOW = datetime.now()
D = [(NOW - timedelta(days=i)).strftime("%Y-%m-%d") for i in range(7)]  # D[0]=today .. D[6]=6 days ago
START, END = D[6], D[0]


def _login():
    for pw in ("9xsecurity", "Temp123456"):
        r = requests.post(f"{BASE}/api/login", json={"username": "admin", "password": pw}, timeout=10)
        if r.status_code == 200:
            return r.json()["token"]
    pytest.skip("login failed")


@pytest.fixture(scope="module")
def tok():
    return _login()


@pytest.fixture(scope="module")
def H(tok):
    return {"X-Auth-Token": tok}


_n = [0]


def _ins(c, day, t, vt, d, plate="", count=1, gate="Main"):
    _n[0] += 1
    p = os.path.join(SNAP, f"{TAG}_{_n[0]}.jpg")
    c.execute("INSERT INTO events(timestamp,date,time,vehicle_type,direction,plate,image_path,plate_status,plate_source,count,gate)"
              " VALUES(?,?,?,?,?,?,?,?,?,?,?)", (f"{day}T{t}", day, t, vt, d, plate, p,
                                                 "done" if plate else "", "cloud" if plate else "", count, gate))


@pytest.fixture(scope="module", autouse=True)
def seed():
    c = sqlite3.connect(DB)
    _ins(c, D[6], "09:30:00", "truck", "Entry", "TQ45AB9876")
    _ins(c, D[6], "18:00:00", "truck", "Exit", "TQ45AB9876")
    _ins(c, D[4], "10:00:00", "car", "Entry", "", gate="Back")          # 4-wheeler without plate
    _ins(c, D[4], "10:05:00", "motorcycle", "Entry", gate="Back")
    _ins(c, D[2], "11:00:00", "person", "Exit", count=2, gate="Back")
    _ins(c, D[0], "07:00:00", "truck", "Entry", "TQ45AB9876")
    _ins(c, D[0], "07:30:00", "car", "Exit", "TQ45CD9876", gate="Back")
    c.commit()
    yield
    c.execute("DELETE FROM events WHERE image_path LIKE ?", (f"%{TAG}%",))
    c.commit()
    c.close()
    for f in os.listdir(report.REPORT_DIR) if os.path.isdir(report.REPORT_DIR) else []:
        if f.startswith(f"9x_Report_{START}_to_") or f.startswith("9x_Report_2026-09-2") or f.startswith(f"9x_Report_{D[0]}_to_"):
            try:
                os.remove(os.path.join(report.REPORT_DIR, f))
            except OSError:
                pass


# ---- range JSON
def test_range_summary_and_days(H):
    r = requests.get(f"{BASE}/api/report/range", params={"start": START, "end": END, "fmt": "pdf"}, headers=H, timeout=60)
    assert r.status_code == 200, r.text
    j = r.json()
    s = j["summary"]
    assert s["events"] == 7 and s["entry"] == 4 and s["exit"] == 4  # person x2
    assert s["start_day"] == START and s["end_day"] == END
    assert s["total"]["truck"] == {"Entry": 2, "Exit": 1}
    assert s["total"]["two_wheeler"]["Entry"] == 1
    assert s["gates"]["Back"]["person"]["Exit"] == 2
    days = j["days"]
    assert [d["date"] for d in days] == list(reversed(D))
    assert days[0]["events"] == 2 and days[1]["events"] == 0 and days[1]["entry"] == 0 and days[2]["events"] == 2
    assert days[4]["exit"] == 2 and days[6]["entry"] == 1 and days[6]["exit"] == 1
    assert sum(d["events"] for d in days) == 7
    assert "se" in j["period"] and "(7 din)" in j["period"]
    assert j["path"].endswith(f"9x_Report_{START}_to_{END}.pdf") and os.path.exists(j["path"])
    assert j["pdf_url"] == f"/api/report/range.pdf?start={START}&end={END}"
    assert j["xlsx_url"] == f"/api/report/range.xlsx?start={START}&end={END}"


def test_range_xlsx_fmt_path(H):
    j = requests.get(f"{BASE}/api/report/range", params={"start": START, "end": END, "fmt": "xlsx"}, headers=H, timeout=60).json()
    assert j["path"].endswith(".xlsx") and os.path.exists(j["path"])


def test_range_single_day_period(H):
    j = requests.get(f"{BASE}/api/report/range", params={"start": D[0], "end": D[0]}, headers=H, timeout=60).json()
    assert len(j["days"]) == 1 and j["period"].endswith("(poora din)") and j["summary"]["events"] == 2


@pytest.mark.parametrize("start,end,msg", [
    (D[0], D[6], "'se' date 'tak' date ke baad hai"),
    ("2024-01-01", "2026-01-01", "1 saal"),
    ("2026-13-01", "2026-10-01", "YYYY-MM-DD"),
    ("abc", "2026-10-01", "YYYY-MM-DD"),
])
def test_range_validation(H, start, end, msg):
    r = requests.get(f"{BASE}/api/report/range", params={"start": start, "end": end}, headers=H, timeout=30)
    assert r.status_code == 400 and msg in r.json()["detail"], r.text
    r = requests.get(f"{BASE}/api/report/range.pdf", params={"start": start, "end": end}, headers=H, timeout=30)
    assert r.status_code == 400


def test_range_requires_auth():
    assert requests.get(f"{BASE}/api/report/range", params={"start": START, "end": END}, timeout=10).status_code == 401
    assert requests.get(f"{BASE}/api/report/range.pdf", params={"start": START, "end": END, "t": "bad"}, timeout=10).status_code == 401


def test_range_pdf_download(tok):
    r = requests.get(f"{BASE}/api/report/range.pdf", params={"start": START, "end": END, "t": tok}, timeout=60)
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/pdf") and r.content[:5] == b"%PDF-"
    import fitz

    text = "\n".join(p.get_text() for p in fitz.open(stream=r.content, filetype="pdf"))
    assert "9x Security - Report" in text and "Day-wise (7 days)" in text and "All events (7)" in text
    assert "TQ45AB9876" in text and "TQ45CD9876" in text and "Not detected" in text
    dd = datetime.strptime(D[6], "%Y-%m-%d").strftime("%d-%m")
    assert f"{dd} 09:30 AM" in text


def test_range_pdf_single_day_daily(tok):
    r = requests.get(f"{BASE}/api/report/range.pdf", params={"start": D[0], "end": D[0], "t": tok}, timeout=60)
    import fitz

    assert "Daily Report" in fitz.open(stream=r.content, filetype="pdf")[0].get_text()


def test_range_xlsx_download(tok):
    r = requests.get(f"{BASE}/api/report/range.xlsx", params={"start": START, "end": END, "t": tok}, timeout=60)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Day-wise", "Events"]
    rows = [[c.value for c in row] for row in wb["Summary"].iter_rows()]
    assert ["Gate", "Type", "Entry", "Exit", "Total"] in rows
    assert any(r_ and r_[0] == "All gates" for r_ in rows)
    assert ["Main", "TOTAL", 2, 1, 3] in rows and ["Back", "TOTAL", 2, 3, 5] in rows
    wd = wb["Day-wise"]
    assert wd.max_row == 8 and wd.max_column == 17
    we = wb["Events"]
    assert [c.value for c in we[1]] == ["Date", "Time", "Gate", "Type", "Direction", "Qty", "Plate", "Plate source", "Photo"]
    assert we.max_row == 8 and we.freeze_panes == "A2"
    # column widths bounded
    for ws in wb.worksheets:
        for k, dim in ws.column_dimensions.items():
            if dim.width:
                assert 10 <= dim.width <= 48, (ws.title, k, dim.width)


def test_range_empty_pdf(tok):
    r = requests.get(f"{BASE}/api/report/range.pdf", params={"start": "2026-09-20", "end": "2026-09-22", "t": tok}, timeout=60)
    import fitz

    assert r.status_code == 200 and "No Entry/Exit recorded in this period" in fitz.open(stream=r.content, filetype="pdf")[0].get_text()
    r = requests.get(f"{BASE}/api/report/range.xlsx", params={"start": "2026-09-20", "end": "2026-09-22", "t": tok}, timeout=60)
    assert r.status_code == 200 and r.content[:2] == b"PK"


# ---- plate search
def test_search_partial_and_order(H):
    j = requests.get(f"{BASE}/api/plate/search", params={"q": "9876"}, headers=H, timeout=20).json()
    assert j["q"] == "9876" and not j["fuzzy"] and len(j["events"]) == 4
    assert j["plates"] == ["TQ45AB9876", "TQ45CD9876"] and j["entries"] == 2 and j["exits"] == 2
    ts = [e["timestamp"] for e in j["events"]]
    assert ts == sorted(ts, reverse=True), "events must be newest first"
    assert j["last_seen"] == ts[0] and j["first_seen"] == ts[-1]


def test_search_normalises(H):
    j = requests.get(f"{BASE}/api/plate/search", params={"q": "tq-45 ab 9876"}, headers=H, timeout=20).json()
    assert j["q"] == "TQ45AB9876" and len(j["events"]) == 3 and j["plates"] == ["TQ45AB9876"]


def test_search_fuzzy(H):
    j = requests.get(f"{BASE}/api/plate/search", params={"q": "TQ459876"}, headers=H, timeout=20).json()
    assert j["fuzzy"] is True and set(j["plates"]) == {"TQ45AB9876", "TQ45CD9876"}
    ts = [e["timestamp"] for e in j["events"]]
    assert ts == sorted(ts, reverse=True)


def test_search_none_and_short(H):
    j = requests.get(f"{BASE}/api/plate/search", params={"q": "ZZZZ"}, headers=H, timeout=20).json()
    assert j["events"] == [] and j["fuzzy"] is False and j["first_seen"] == ""
    for q in ("", "a", "-1-", "  "):
        r = requests.get(f"{BASE}/api/plate/search", params={"q": q}, headers=H, timeout=20)
        assert r.status_code == 400 and "Kam se kam 2" in r.json()["detail"]
    assert requests.get(f"{BASE}/api/plate/search", params={"q": "1234"}, timeout=10).status_code == 401


def test_search_like_wildcards_not_injected(H):
    # '%' and '_' are stripped by normalisation -> cannot match everything
    r = requests.get(f"{BASE}/api/plate/search", params={"q": "%_%"}, headers=H, timeout=20)
    assert r.status_code == 400


# ---- Plate Recognizer usage (FAKE key)
def test_usage_no_key_then_fake_key(H):
    orig = requests.get(f"{BASE}/api/settings", headers=H, timeout=10).json()
    try:
        requests.post(f"{BASE}/api/settings", json={"pr_api_key": "", "pr_enabled": False}, headers=H, timeout=10)
        j = requests.get(f"{BASE}/api/plate/cloud/usage", headers=H, timeout=20).json()
        assert j == {"enabled": False, "key_set": False, "valid": False, "message": "API key nahi daali"}
        assert requests.post(f"{BASE}/api/settings", json={"pr_api_key": "FAKEKEY123", "pr_enabled": True}, headers=H, timeout=10).status_code == 200
        t0 = time.time()
        j = requests.get(f"{BASE}/api/plate/cloud/usage?refresh=1", headers=H, timeout=30).json()
        t_first = time.time() - t0
        assert j["enabled"] and j["key_set"] and j["valid"] is False and j["message"].startswith("HTTP 403: API key galat hai"), j
        assert j["checked_at"]
        t0 = time.time()
        j2 = requests.get(f"{BASE}/api/plate/cloud/usage", headers=H, timeout=30).json()
        t_cached = time.time() - t0
        assert j2["checked_at"] == j["checked_at"], "second call should be served from cache"
        assert t_cached < max(0.3, t_first)
        # disabled but key set -> enabled false, still cached result
        requests.post(f"{BASE}/api/settings", json={"pr_enabled": False}, headers=H, timeout=10)
        j3 = requests.get(f"{BASE}/api/plate/cloud/usage", headers=H, timeout=30).json()
        assert j3["enabled"] is False and j3["key_set"] is True and j3["checked_at"] == j["checked_at"]
        # different key -> cache miss
        requests.post(f"{BASE}/api/settings", json={"pr_api_key": "FAKEKEY456"}, headers=H, timeout=10)
        time.sleep(1.1)
        j4 = requests.get(f"{BASE}/api/plate/cloud/usage", headers=H, timeout=30).json()
        assert j4["checked_at"] != j["checked_at"]
    finally:
        requests.post(f"{BASE}/api/settings", json={"pr_api_key": orig.get("pr_api_key", "") or "",
                                                      "pr_enabled": bool(orig.get("pr_enabled"))}, headers=H, timeout=10)


# ---- pure-unit edges
def test_range_window_boundaries():
    s, e = report.range_window("2026-01-31", "2026-02-01")  # month boundary
    assert (e - s).days == 2
    report.range_window("2025-01-01", "2025-12-31")  # 365 calendar days
    report.range_window("2024-01-01", "2024-12-31")  # 366 calendar days (leap) - must pass
    with pytest.raises(ValueError):
        report.range_window("2025-01-01", "2026-01-03")  # 368 calendar days


def test_range_window_367_days_limit():
    """Message says max 366 din; 2025-01-01..2026-01-02 is 367 calendar days."""
    with pytest.raises(ValueError):
        report.range_window("2025-01-01", "2026-01-02")


def test_month_boundary_days(tmp_path, monkeypatch):
    monkeypatch.setattr(report, "REPORT_DIR", str(tmp_path))
    d = database.EventDB(db_path=str(tmp_path / "e.db"))
    d.add_event("truck", "Entry", "AB12CD3456", "/s/a.jpg", ts=datetime(2026, 1, 31, 23, 59, 59), gate="G")
    d.add_event("truck", "Exit", "AB12CD3456", "/s/b.jpg", ts=datetime(2026, 2, 1, 0, 0, 0), gate="G")
    d.add_event("truck", "Exit", "AB12CD3456", "/s/c.jpg", ts=datetime(2026, 2, 2, 0, 0, 0), gate="G")  # outside
    summary, days, path = report.generate_range(d, "2026-01-31", "2026-02-01", "pdf")
    assert summary["events"] == 2 and [x["date"] for x in days] == ["2026-01-31", "2026-02-01"]
    assert days[0]["entry"] == 1 and days[1]["exit"] == 1
    assert report.period_text(summary) == "31-01-2026 se 01-02-2026 tak (2 din)"
